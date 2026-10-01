"""
Módulo de Exportación de Resultados
Genera reportes formateados en Excel (.xlsx) y empaqueta capas de error espacial en formato ZIP.
"""

import os
import io
import zipfile
import tempfile
import pandas as pd
import geopandas as gpd
from typing import Dict, List, Optional
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

def generate_excel_report(report_df: pd.DataFrame, spatial_logs: List[Dict]) -> bytes:
    """
    Genera un archivo Excel en memoria con formato profesional y colores de alerta.
    Pestaña 1: Resumen de Diagnóstico Espacial y Log.
    Pestaña 2: Detalle de Verificación de Campos y Atributos.
    """
    output = io.BytesIO()
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        # Pestaña 1: Logs Espaciales
        logs_df = pd.DataFrame(spatial_logs)
        if logs_df.empty:
            logs_df = pd.DataFrame([{"Mensaje": "Sin errores espaciales registrados", "Tipo": "OK"}])
            
        logs_df.to_excel(writer, sheet_name="Resumen_Espacial", index=False)
        
        # Pestaña 2: Atributos
        if not report_df.empty:
            report_df.to_excel(writer, sheet_name="Revision_Campos", index=False)
        else:
            pd.DataFrame([{"Mensaje": "Sin datos de revisión"}]).to_excel(writer, sheet_name="Revision_Campos", index=False)

    # Estilizar con OpenPyXL
    output.seek(0)
    wb = openpyxl.load_workbook(output)

    # Estilos
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    
    green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    
    green_font = Font(color="375623", bold=True)
    red_font = Font(color="C65911", bold=True)

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True
        
        # Formatear Encabezados
        for col_num, col_name in enumerate(ws.iter_cols(min_row=1, max_row=1), 1):
            for cell in col_name:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Formatear Celda a Celda y ajustar anchos
        for row in ws.iter_rows(min_row=2, max_col=ws.max_column, max_row=ws.max_row):
            for cell in row:
                cell.border = thin_border
                val_str = str(cell.value)
                if val_str in ["Sí", "True", "TRUE", "OK", "CORRECTO"]:
                    cell.fill = green_fill
                    cell.font = green_font
                elif val_str in ["No", "False", "FALSE", "ERROR", "INCORRECTO"]:
                    cell.fill = red_fill
                    cell.font = red_font

        # Autoajustar ancho de columnas
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    final_output = io.BytesIO()
    wb.save(final_output)
    return final_output.getvalue()

def export_spatial_errors_to_zip(spatial_errors: Dict[str, gpd.GeoDataFrame]) -> Optional[bytes]:
    """
    Empaqueta los GeoDataFrames de geometrías conflictivas/errores en un archivo ZIP descargable.
    """
    if not spatial_errors:
        return None

    zip_buffer = io.BytesIO()

    with tempfile.TemporaryDirectory() as temp_dir:
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            for layer_name, error_gdf in spatial_errors.items():
                if error_gdf is not None and not error_gdf.empty:
                    shp_base = f"ERROR_{layer_name}"
                    shp_path = os.path.join(temp_dir, f"{shp_base}.shp")
                    
                    # Guardar a shapefile en temp_dir
                    error_gdf.to_file(shp_path, driver="ESRI Shapefile")

                    # Agregar todos los componentes del shapefile (.shp, .shx, .dbf, .prj) al ZIP
                    for ext in [".shp", ".shx", ".dbf", ".prj", ".cpg"]:
                        comp_path = os.path.join(temp_dir, f"{shp_base}{ext}")
                        if os.path.exists(comp_path):
                            zip_file.write(comp_path, arcname=f"{shp_base}{ext}")

    zip_buffer.seek(0)
    return zip_buffer.getvalue()
